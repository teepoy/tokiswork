"""Core PPTX agent logic.

This module contains the reusable implementation layers for PPTX workflows:
- template parsing
- agent orchestration
- template creation
- high-level run helpers
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import dspy
from pptx import Presentation
from pptx.util import Inches, Pt


@dataclass
class TemplateField:
    """Represents a fillable field in a PPTX template."""

    slide_idx: int
    shape_idx: int
    placeholder_text: str
    field_name: str
    is_required: bool
    default_value: str | None = None


@dataclass
class PPTXAgentResult:
    """Result of PPTX agent execution."""

    output_path: Path
    filled_fields: dict[str, str]
    unfilled_fields: list[str]
    metadata: dict[str, Any]


class AnalyzeTemplate(dspy.Signature):
    """Analyze PPTX template structure and extract fillable regions."""

    template_info: str = dspy.InputField(
        desc="JSON info about template structure including placeholders and field names"
    )
    analysis: str = dspy.OutputField(
        desc="Detailed analysis of template fields, requirements, and types"
    )
    field_names: str = dspy.OutputField(
        desc="Comma-separated list of field names to extract from user"
    )


class ExtractInformation(dspy.Signature):
    """Extract relevant information from user input based on template requirements."""

    field_requirements: str = dspy.InputField(
        desc="JSON describing required fields and their requirements"
    )
    user_input: str = dspy.InputField(desc="User-provided input/context")
    extracted_data: str = dspy.OutputField(
        desc="JSON object with extracted values for each field"
    )
    missing_fields: str = dspy.OutputField(
        desc="Comma-separated list of fields not found in user input"
    )
    reasoning: str = dspy.OutputField(
        desc="Explanation of extraction decisions and any assumptions made"
    )


class GenerateUserPrompt(dspy.Signature):
    """Generate a friendly prompt asking user for missing information."""

    missing_fields: str = dspy.InputField(
        desc="Comma-separated list of field names needed from user"
    )
    field_descriptions: str = dspy.InputField(
        desc="JSON object describing what each field represents"
    )
    user_prompt: str = dspy.OutputField(
        desc="Clear question/prompt to ask user for the missing information"
    )


class ProcessUserResponse(dspy.Signature):
    """Process user's response to a data request."""

    user_response: str = dspy.InputField(
        desc="User's response providing missing information"
    )
    expected_fields: str = dspy.InputField(
        desc="Comma-separated list of expected field names"
    )
    extracted_values: str = dspy.OutputField(
        desc="JSON object mapping field names to extracted values"
    )
    completeness: str = dspy.OutputField(
        desc="Assessment of whether user provided sufficient information"
    )


class PPTXTemplateParser:
    """Parse and extract information from PPTX template files."""

    PLACEHOLDER_PATTERN = r"\{\{(\w+)\}\}|\[([A-Z_]+)\]|<<(\w+)>>"

    def __init__(self, template_path: str | Path):
        self.template_path = Path(template_path)
        self.presentation = Presentation(str(self.template_path))
        self.fields: list[TemplateField] = []
        self._parse_template()

    def _parse_template(self) -> None:
        for slide_idx, slide in enumerate(self.presentation.slides):
            for shape_idx, shape in enumerate(slide.shapes):
                if not hasattr(shape, "text"):
                    continue

                text = shape.text
                matches = re.finditer(self.PLACEHOLDER_PATTERN, text)
                for match in matches:
                    field_name = match.group(1) or match.group(2) or match.group(3)
                    field_name_lower = field_name.lower()
                    is_required = not any(
                        marker in text.lower()
                        for marker in ["optional", "default", "if applicable"]
                    )
                    default_value = self._extract_default_value(text, field_name)
                    self.fields.append(
                        TemplateField(
                            slide_idx=slide_idx,
                            shape_idx=shape_idx,
                            placeholder_text=match.group(0),
                            field_name=field_name_lower,
                            is_required=is_required,
                            default_value=default_value,
                        )
                    )

    def _extract_default_value(self, text: str, field_name: str) -> str | None:
        default_pattern = (
            rf'{re.escape(field_name)}\s*:?\s*(?:default|=)\s*["\']?([^"\']*)["\']?'
        )
        match = re.search(default_pattern, text, re.IGNORECASE)
        return match.group(1) if match else None

    def get_field_info(self) -> dict[str, Any]:
        return {
            "total_fields": len(self.fields),
            "required_fields": [f.field_name for f in self.fields if f.is_required],
            "optional_fields": [f.field_name for f in self.fields if not f.is_required],
            "fields_with_defaults": {
                f.field_name: f.default_value for f in self.fields if f.default_value
            },
            "all_fields": [f.field_name for f in self.fields],
        }

    def fill_template(
        self, field_values: dict[str, str], output_path: str | Path
    ) -> None:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs = Presentation(str(self.template_path))

        for field in self.fields:
            slide = prs.slides[field.slide_idx]
            shape = slide.shapes[field.shape_idx]
            if not hasattr(shape, "text_frame"):
                continue

            value = field_values.get(
                field.field_name,
                field.default_value or f"[{field.field_name.upper()}]",
            )
            original_text = shape.text
            new_text = original_text.replace(field.placeholder_text, str(value))
            text_frame = shape.text_frame
            if len(text_frame.paragraphs) > 0:
                paragraph = text_frame.paragraphs[0]
                if len(paragraph.runs) > 0:
                    paragraph.runs[0].text = new_text
                else:
                    paragraph.text = new_text


        prs.save(str(output_path))


class PPTXReActAgent(dspy.Module):
    """ReAct agent for PPTX template processing."""

    def __init__(
        self,
        analyze_module: dspy.Module | None = None,
        extract_module: dspy.Module | None = None,
        generate_prompt_module: dspy.Module | None = None,
        process_response_module: dspy.Module | None = None,
    ):
        super().__init__()
        self.analyze = analyze_module or dspy.Predict(AnalyzeTemplate)
        self.extract = extract_module or dspy.Predict(ExtractInformation)
        self.generate_prompt = generate_prompt_module or dspy.Predict(
            GenerateUserPrompt
        )
        self.process_response = process_response_module or dspy.Predict(
            ProcessUserResponse
        )
        self.thought_log: list[dict[str, str]] = []

    def _log_thought(self, action: str, observation: str) -> None:
        self.thought_log.append({"action": action, "observation": observation})

    @staticmethod
    def _safe_json_dict(text: str) -> dict[str, str]:
        if not text.strip().startswith("{"):
            return {}
        try:
            data = json.loads(text)
        except json.JSONDecodeError:
            return {}
        return {str(k): str(v) for k, v in data.items()}

    def forward(
        self,
        template_path: str | Path,
        user_input: str,
        output_path: str | Path,
        user_interaction_callback: Callable[[str], str] | None = None,
    ) -> PPTXAgentResult:
        self.thought_log = []
        output_path = Path(output_path)

        self._log_thought("ANALYZE_TEMPLATE", f"Parsing template file: {template_path}")
        parser = PPTXTemplateParser(template_path)
        field_info = parser.get_field_info()
        self._log_thought(
            "TEMPLATE_ANALYSIS_COMPLETE",
            json.dumps(field_info, ensure_ascii=False, indent=2),
        )

        self._log_thought(
            "EXTRACT_INFORMATION", f"Processing user input: {user_input[:100]}..."
        )
        field_requirements = json.dumps(
            {
                "required": field_info["required_fields"],
                "optional": field_info["optional_fields"],
                "defaults": field_info["fields_with_defaults"],
            },
            ensure_ascii=False,
        )
        extraction_result = self.extract(
            field_requirements=field_requirements,
            user_input=user_input,
        )
        self._log_thought(
            "EXTRACTION_COMPLETE",
            f"Missing fields: {extraction_result.missing_fields}",
        )

        extracted_data = self._safe_json_dict(extraction_result.extracted_data)
        missing_fields = [
            f.strip() for f in extraction_result.missing_fields.split(",") if f.strip()
        ]
        all_data = extracted_data.copy()

        if missing_fields:
            self._log_thought(
                "MISSING_FIELDS_DETECTED",
                f"Need to collect: {', '.join(missing_fields)}",
            )
            required_missing = [
                f for f in missing_fields if f in field_info["required_fields"]
            ]
            if required_missing and user_interaction_callback:
                prompt_result = self.generate_prompt(
                    missing_fields=", ".join(required_missing),
                    field_descriptions=json.dumps(
                        {f: f"Required information for {f}" for f in required_missing},
                        ensure_ascii=False,
                    ),
                )
                self._log_thought(
                    "GENERATED_USER_PROMPT",
                    f"Asking user for: {', '.join(required_missing)}",
                )
                user_response = user_interaction_callback(prompt_result.user_prompt)
                self._log_thought("USER_RESPONSE_RECEIVED", user_response[:200])
                response_result = self.process_response(
                    user_response=user_response,
                    expected_fields=", ".join(required_missing),
                )
                all_data.update(self._safe_json_dict(response_result.extracted_values))

            for field_name in missing_fields:
                if field_name not in all_data:
                    default = field_info["fields_with_defaults"].get(field_name)
                    if default:
                        all_data[field_name] = default
                        self._log_thought("USING_DEFAULT", f"{field_name} = {default}")

        self._log_thought("GENERATE_OUTPUT", f"Creating output PPTX at {output_path}")
        parser.fill_template(all_data, output_path)
        self._log_thought(
            "GENERATION_COMPLETE", f"Successfully generated: {output_path}"
        )

        unfilled = [
            f for f in field_info["all_fields"] if f not in all_data or all_data[f].startswith("[")
        ]
        return PPTXAgentResult(
            output_path=output_path,
            filled_fields=all_data,
            unfilled_fields=unfilled,
            metadata={
                "template_path": str(template_path),
                "template_fields_total": field_info["total_fields"],
                "required_fields": field_info["required_fields"],
                "optional_fields": field_info["optional_fields"],
                "thought_log": self.thought_log,
                "reasoning": extraction_result.reasoning,
            },
        )


def create_simple_template(output_path: str | Path, fields: dict[str, str]) -> None:
    """Create a simple PPTX template file for testing."""

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    title_frame = title_box.text_frame
    title_frame.text = "Document Template"
    title_frame.paragraphs[0].font.size = Pt(44)
    title_frame.paragraphs[0].font.bold = True

    y_position = 2.0
    for field_name, placeholder in fields.items():
        label_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(y_position), Inches(2), Inches(0.4)
        )
        label_frame = label_box.text_frame
        label_frame.text = f"{field_name}:"
        label_frame.paragraphs[0].font.bold = True

        value_box = slide.shapes.add_textbox(
            Inches(2.5), Inches(y_position), Inches(7), Inches(0.4)
        )
        value_frame = value_box.text_frame
        value_frame.text = placeholder
        y_position += 0.6

    prs.save(str(output_path))


def analyze_template(template_path: str | Path) -> dict[str, Any]:
    """Analyze a PPTX template without filling it."""

    parser = PPTXTemplateParser(template_path)
    return {
        "template_path": str(template_path),
        "field_info": parser.get_field_info(),
        "total_slides": len(parser.presentation.slides),
    }


def create_template_from_spec(output_path: str | Path, spec: dict[str, str]) -> str:
    """Create a template PPTX from a field spec."""

    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    placeholders = {name: f"{{{{{name}}}}}" for name in spec.keys()}
    create_simple_template(output_path, placeholders)
    return str(output_path)


def run_pptx_agent(
    template_path: str | Path,
    user_input: str,
    output_dir: str | Path = "pptx_runs",
    output_path: str | Path | None = None,
    use_real_lm: bool = False,
    interactive: bool = False,
) -> dict[str, Any]:
    """Run the PPTX agent and persist metadata alongside the result."""

    if use_real_lm:
        # Hook for external LM configuration by caller/environment.
        pass

    template_path = Path(template_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    final_output_path = Path(output_path).resolve() if output_path else output_dir / f"{template_path.stem}_output.pptx"

    agent = PPTXReActAgent()

    def interaction_callback(prompt: str) -> str:
        if interactive:
            print("\n" + "=" * 60)
            print(prompt)
            print("=" * 60)
            return input("Please provide the information: ").strip()
        return "Using default values for optional fields."

    result = agent.forward(
        template_path=template_path,
        user_input=user_input,
        output_path=final_output_path,
        user_interaction_callback=interaction_callback if interactive else None,
    )

    result_dict = {
        "success": final_output_path.exists(),
        "output_path": str(final_output_path),
        "filled_fields": result.filled_fields,
        "unfilled_fields": result.unfilled_fields,
        "metadata": result.metadata,
    }
    metadata_path = output_dir / f"{template_path.stem}_metadata.json"
    metadata_path.write_text(
        json.dumps(result_dict, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    result_dict["metadata_path"] = str(metadata_path)
    return result_dict
